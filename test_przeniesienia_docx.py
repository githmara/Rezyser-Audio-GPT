"""
test_przeniesienia_docx.py — przeniesienia wyrazów w `.docx` z konwersji PDF→Word (19.15).

Zmierzone na `.docx` słownika koronawirusowego (2026-10-20): 5 890 miejsc
„по- священного” w środku akapitu. `core_poliglota.napraw_przeniesienia`
rozstrzyga WYŁĄCZNIE słownikiem tego samego dokumentu:
  1. forma sklejona występuje w dokumencie → sklejamy;
  2. forma z łącznikiem występuje częściej → łącznik zostaje, spacja znika;
  3. brak dowodu albo remis → bez zmian: dywizy wiszące („аудио- и
     видеоконференций”), slugi adresów, przeniesienie przez granicę akapitu.
Obie ścieżki `.docx` (Poliglota `tekst_docx`, Konwerter) idą tą funkcją.

Uruchom:  .venv/Scripts/python -m pytest test_przeniesienia_docx.py -q
Albo jako skrypt (deleguje do pytesta): .venv/Scripts/python test_przeniesienia_docx.py
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))

import core_poliglota as cp  # noqa: E402

# Słownik dokumentu: słowa, które występują gdzie indziej BEZ przeniesienia.
SLOWNIK = ("Ковид и посвященного. Ковид-диссидентов много, по-прежнему. "
           "Видеоконференций тоже.")


def _napraw(*akapity):
    return cp.napraw_przeniesienia([SLOWNIK, *akapity])


def test_sklejone_gdy_forma_sklejona_jest_w_dokumencie():
    wynik, sklejone, laczniki = _napraw("Издание, по- священного теме.")
    assert wynik[1] == "Издание, посвященного теме."
    assert (sklejone, laczniki) == (1, 0)


def test_lacznik_gdy_zlozenie_zlamane_na_swoim_laczniku():
    wynik, sklejone, laczniki = _napraw("Мнение ковид- диссидентов и по- прежнему.")
    assert wynik[1] == "Мнение ковид-диссидентов и по-прежнему."
    assert (sklejone, laczniki) == (0, 2)


@pytest.mark.parametrize("akapit", [
    "Сеансы аудио- и видеоконференций.",           # dywiz wiszący
    "Пред- и постпандемийный период.",              # dywiz wiszący, wielka litera
    "Источник: obozrenie- birulevo, 2020.",         # slug adresu
    "Кое-что и по-английски.",                      # łącznik bez spacji
    "Слово - слово – слово — слово.",               # myślniki
    "Вирус COVID- 19 и Ковид- Москва.",             # prawa część nie od małej
    "Нет доказательства: встре- чаются.",           # brak słowa w dokumencie
])
def test_bez_dowodu_bez_zmian(akapit):
    wynik, sklejone, laczniki = _napraw(akapit)
    assert wynik[1] == akapit and (sklejone, laczniki) == (0, 0)


def test_liczba_akapitow_i_granica_akapitu_bez_zmian():
    # Naprawiacz tagów nakłada kody PO INDEKSIE akapitu oryginału, a
    # przeniesienie przez granicę akapitu zmieniałoby podział.
    akapity = [SLOWNIK, "Конец строки по-", "священного начало.", ""]
    wynik, sklejone, _ = cp.napraw_przeniesienia(akapity)
    assert wynik == akapity and sklejone == 0


def _docx(sciezka, *akapity):
    import docx
    d = docx.Document()
    for a in akapity:
        d.add_paragraph(a)
    d.save(str(sciezka))


def test_tekst_docx_poligloty_skleja(tmp_path):
    p = tmp_path / "slownik.docx"
    _docx(p, SLOWNIK, "Издание, по- священного теме.", "Сеансы аудио- и видеоконференций.")
    linie = cp.tekst_docx(str(p)).split("\n")
    assert linie[1:] == ["Издание, посвященного теме.", "Сеансы аудио- и видеоконференций."]


def test_konwerter_skleja_tak_samo(tmp_path, monkeypatch):
    wx = pytest.importorskip("wx")
    import docx
    import gui_konwerter as gk
    import i18n
    p = tmp_path / "slownik.docx"
    _docx(p, SLOWNIK, "Издание, по- священного теме.")
    zapisane = []
    monkeypatch.setattr(docx.document.Document, "save",
                        lambda self, _p: zapisane.append([x.text for x in self.paragraphs]))
    monkeypatch.setattr(wx, "MessageBox", lambda *_a, **_k: wx.OK)
    app = wx.App(False)
    jezyk_przed = i18n.aktualny_jezyk()
    i18n.ustaw_jezyk("pl")
    ramka = wx.Frame(None)
    try:
        panel = gk.KonwerterPanel(ramka)
        panel._txt_file.SetValue(str(p))
        panel._on_build(None)
    finally:
        ramka.Destroy()
        app.Destroy()
        i18n.ustaw_jezyk(jezyk_przed)
    assert zapisane, "Konwerter nic nie zapisał"
    tekst = "\n".join(zapisane[-1])
    assert "посвященного" in tekst and "по- священного" not in tekst


# Uruchomienie jako skrypt — deleguje do pytesta (kanon v19.2.1).
if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
